import io
import re
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def _add_styled_runs(paragraph, text: str) -> None:
    """Parses basic inline Markdown (*italic*, **bold**, ***bold italic***, `code`) into docx runs."""
    pattern = r"(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)"
    tokens = re.split(pattern, text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("***") and token.endswith("***") and len(token) >= 6:
            run = paragraph.add_run(token[3:-3])
            run.bold = True
            run.italic = True
        elif token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(99, 102, 241)
        else:
            paragraph.add_run(token)


def markdown_to_docx(markdown_content: str, title: Optional[str] = None) -> io.BytesIO:
    """Converts markdown formatted summary into a beautifully styled Word (.docx) document.
    
    Returns an in-memory BytesIO stream ready for download.
    """
    doc = Document()
    
    # Configure 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Add Document Title if provided
    if title:
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(title)
        title_run.font.name = "Arial"
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(30, 41, 59)  # Slate 800
        title_p.paragraph_format.space_after = Pt(12)
        title_p.paragraph_format.space_before = Pt(0)
    
    lines = markdown_content.splitlines()
    in_code_block = False
    code_block_lines = []
    
    for line in lines:
        stripped = line.strip()
        
        # Handle code blocks
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(6)
                code_text = "\n".join(code_block_lines)
                run = p.add_run(code_text)
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(71, 85, 105)
                code_block_lines = []
            else:
                in_code_block = True
            continue
        
        if in_code_block:
            code_block_lines.append(line)
            continue
        
        # Blank lines
        if not stripped:
            continue
        
        # Headings
        if stripped.startswith("# "):
            p = doc.add_heading(level=1)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            _add_styled_runs(p, stripped[2:].strip())
        elif stripped.startswith("## "):
            p = doc.add_heading(level=2)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            _add_styled_runs(p, stripped[3:].strip())
        elif stripped.startswith("### "):
            p = doc.add_heading(level=3)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            _add_styled_runs(p, stripped[4:].strip())
        elif stripped.startswith("#### "):
            p = doc.add_heading(level=4)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            _add_styled_runs(p, stripped[5:].strip())
        # Blockquotes
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(4)
            _add_styled_runs(p, stripped[2:].strip())
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(100, 116, 139)
        # Bullet list items
        elif re.match(r"^[-*]\s+", stripped):
            item_text = re.sub(r"^[-*]\s+", "", stripped)
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            _add_styled_runs(p, item_text)
        # Numbered list items
        elif re.match(r"^\d+\.\s+", stripped):
            item_text = re.sub(r"^\d+\.\s+", "", stripped)
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            _add_styled_runs(p, item_text)
        # Standard paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            _add_styled_runs(p, stripped)
    
    # Save to BytesIO buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
